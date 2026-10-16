"""SQLite store for CVPR 2024 paper rows."""

from __future__ import annotations

import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / "data" / "papers.db"
XLSX_PATH = ROOT / "repo" / "cvpr2024_papers.xlsx"
PAGE_SIZE = 5

_SCHEMA = """
CREATE TABLE IF NOT EXISTS papers (
    paper_url TEXT PRIMARY KEY,
    title TEXT NOT NULL DEFAULT '',
    authors TEXT NOT NULL DEFAULT '',
    abstract TEXT NOT NULL DEFAULT '',
    pdf_url TEXT NOT NULL DEFAULT '',
    supp_url TEXT NOT NULL DEFAULT '',
    arxiv_url TEXT NOT NULL DEFAULT '',
    pages TEXT NOT NULL DEFAULT '',
    abstract_source TEXT NOT NULL DEFAULT '',
    position INTEGER NOT NULL,
    updated_at TEXT NOT NULL
)
"""


_JOB_SCHEMA = """
CREATE TABLE IF NOT EXISTS refresh_job (
    id INTEGER PRIMARY KEY CHECK (id = 1),
    state TEXT NOT NULL,
    message TEXT NOT NULL DEFAULT '',
    total INTEGER
)
"""


def connect() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DB_PATH, timeout=30)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute("PRAGMA busy_timeout=30000")
    connection.execute(_SCHEMA)
    connection.execute(_JOB_SCHEMA)
    connection.execute(
        "INSERT INTO refresh_job (id, state, message, total) VALUES (1, 'idle', '', NULL) "
        "ON CONFLICT(id) DO NOTHING"
    )
    return connection


def count_papers(connection: sqlite3.Connection | None = None) -> int:
    owns_connection = connection is None
    connection = connection or connect()
    try:
        return int(connection.execute("SELECT COUNT(*) FROM papers").fetchone()[0])
    finally:
        if owns_connection:
            connection.close()


def upsert_papers(papers: list[dict[str, str]], *, keep_existing_abstract: bool) -> int:
    if not papers:
        return 0
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    with connect() as connection:
        for position, paper in enumerate(papers, start=1):
            paper_url = (paper.get("paper_url") or "").strip()
            if not paper_url:
                continue
            existing = connection.execute(
                "SELECT abstract, abstract_source, arxiv_url FROM papers WHERE paper_url = ?",
                (paper_url,),
            ).fetchone()
            abstract = (paper.get("abstract") or "").strip()
            abstract_source = (paper.get("abstract_source") or "").strip()
            arxiv_url = (paper.get("arxiv_url") or "").strip()
            if keep_existing_abstract and existing is not None:
                if not abstract:
                    abstract = existing["abstract"]
                    abstract_source = existing["abstract_source"]
                if not arxiv_url:
                    arxiv_url = existing["arxiv_url"]
            connection.execute(
                """
                INSERT INTO papers (
                    paper_url, title, authors, abstract, pdf_url, supp_url,
                    arxiv_url, pages, abstract_source, position, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(paper_url) DO UPDATE SET
                    title = excluded.title,
                    authors = excluded.authors,
                    abstract = excluded.abstract,
                    pdf_url = excluded.pdf_url,
                    supp_url = excluded.supp_url,
                    arxiv_url = excluded.arxiv_url,
                    pages = excluded.pages,
                    abstract_source = excluded.abstract_source,
                    position = excluded.position,
                    updated_at = excluded.updated_at
                """,
                (
                    paper_url,
                    paper.get("title") or "",
                    paper.get("authors") or "",
                    abstract,
                    paper.get("pdf_url") or "",
                    paper.get("supp_url") or "",
                    arxiv_url,
                    str(paper.get("pages") or ""),
                    abstract_source,
                    position,
                    now,
                ),
            )
        connection.commit()
        return count_papers(connection)


def seed_from_xlsx_if_empty() -> None:
    if count_papers() or not XLSX_PATH.is_file():
        return
    workbook = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        papers: list[dict[str, str]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[3]:
                continue
            papers.append(
                {
                    "title": str(row[0] or ""),
                    "authors": str(row[1] or ""),
                    "abstract": str(row[2] or ""),
                    "paper_url": str(row[3] or ""),
                    "pdf_url": str(row[4] or ""),
                    "supp_url": str(row[5] or ""),
                    "arxiv_url": str(row[6] or ""),
                    "pages": str(row[7] or ""),
                    "abstract_source": str(row[8] or ""),
                }
            )
    finally:
        workbook.close()
    upsert_papers(papers, keep_existing_abstract=False)


def papers_missing_abstract() -> list[dict[str, str]]:
    with connect() as connection:
        rows = connection.execute(
            """
            SELECT title, authors, paper_url, pdf_url, supp_url, arxiv_url, pages
            FROM papers
            WHERE abstract = ''
            ORDER BY position
            """
        ).fetchall()
    return [
        {
            "title": row["title"],
            "authors": row["authors"],
            "abstract": "",
            "abstract_source": "",
            "paper_url": row["paper_url"],
            "listing_url": row["paper_url"],
            "pdf_url": row["pdf_url"],
            "supp_url": row["supp_url"],
            "arxiv_url": row["arxiv_url"],
            "pages": row["pages"],
        }
        for row in rows
    ]


def save_fetched_paper(original_url: str, paper: dict[str, str]) -> None:
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    pdf_url = paper.get("pdf_url") or ""
    arxiv_url = paper.get("arxiv_url") or ""
    with connect() as connection:
        connection.execute(
            """
            UPDATE papers SET
                abstract = ?,
                abstract_source = ?,
                pdf_url = CASE WHEN ? != '' THEN ? ELSE pdf_url END,
                arxiv_url = CASE WHEN ? != '' THEN ? ELSE arxiv_url END,
                updated_at = ?
            WHERE paper_url = ?
            """,
            (
                paper.get("abstract") or "",
                paper.get("abstract_source") or "",
                pdf_url,
                pdf_url,
                arxiv_url,
                arxiv_url,
                now,
                original_url,
            ),
        )
        connection.commit()


def get_refresh_job() -> dict[str, object]:
    with connect() as connection:
        row = connection.execute("SELECT state, message, total FROM refresh_job WHERE id = 1").fetchone()
    return {"state": row["state"], "message": row["message"], "total": row["total"]}


def set_refresh_job(state: str, message: str = "", total: int | None = None) -> None:
    with connect() as connection:
        connection.execute(
            "UPDATE refresh_job SET state = ?, message = ?, total = ? WHERE id = 1",
            (state, message, total),
        )
        connection.commit()


def mark_refresh_interrupted() -> None:
    if get_refresh_job()["state"] == "running":
        set_refresh_job("error", "读取中断，请再试一次")


def list_papers(offset: int) -> list[sqlite3.Row]:
    with connect() as connection:
        return list(
            connection.execute(
                """
                SELECT title, authors, abstract, paper_url, pdf_url, pages, position
                FROM papers
                ORDER BY position, paper_url
                LIMIT ? OFFSET ?
                """,
                (PAGE_SIZE, max(offset, 0)),
            )
        )
